"""从电子书 / 办公文档抽出纯文本，供站内阅读与切章。

PDF 不走这里：浏览器原生查看器比抽文本更适合排版页。
抽文本失败时由调用方退回「整文件下载」。
"""
from __future__ import annotations

import io
import struct
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from typing import Dict, List, Optional

# 可抽出正文、走与 txt 相同的站内阅读/切章链路
EXTRACTABLE_EXTS = {
    "txt", "md", "markdown",
    "html", "htm",
    "epub",
    "docx", "doc",
    "rtf",
    "odt",
    "fb2",
}

MAX_EXTRACT_BYTES = 80 * 1024 * 1024  # 过大则不抽，避免把阅读接口打爆


@dataclass
class ExtractedDoc:
    text: str
    parts: List[Dict]  # {title, start, end} 与 split_text 同形


class ExtractError(Exception):
    """无法从该文件抽出可读文本。"""


def extract_document(data: bytes, ext: str) -> ExtractedDoc:
    from .chapters import decode_bytes, split_text

    ext = (ext or "").lower()
    if not data:
        raise ExtractError("文件为空")
    if len(data) > MAX_EXTRACT_BYTES:
        raise ExtractError("文件过大，无法在站内展开阅读，请下载后本地打开")
    if ext in {"txt", "md", "markdown"}:
        text = decode_bytes(data)
        return ExtractedDoc(text=text, parts=split_text(text))
    if ext in {"html", "htm"}:
        text = _html_to_text(decode_bytes(data))
        return ExtractedDoc(text=text, parts=split_text(text))
    if ext == "epub":
        return _extract_epub(data)
    if ext == "docx":
        return _extract_docx(data)
    if ext == "doc":
        return _extract_doc(data)
    if ext == "rtf":
        return _extract_rtf(data)
    if ext == "odt":
        return _extract_odt(data)
    if ext == "fb2":
        return _extract_fb2(data)
    raise ExtractError(f"不支持抽出 .{ext}")


def _html_to_text(html: str) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    text = soup.get_text("\n")
    lines = [ln.strip() for ln in text.splitlines()]
    # 压缩多余空行，保留段落分隔
    out: List[str] = []
    blank = 0
    for ln in lines:
        if not ln:
            blank += 1
            if blank <= 1 and out:
                out.append("")
            continue
        blank = 0
        out.append(ln)
    return "\n".join(out).strip()


def _parts_from_chunks(chunks: List[tuple[Optional[str], str]]) -> ExtractedDoc:
    """chunks: (title, body)。多段则按段成章；否则再对全文做标题切章。"""
    from .chapters import split_text

    nonempty = [(t, b.strip()) for t, b in chunks if (b or "").strip()]
    if not nonempty:
        raise ExtractError("未抽出正文")
    if len(nonempty) == 1:
        title, body = nonempty[0]
        text = f"{title}\n\n{body}" if title and title not in body[:80] else body
        return ExtractedDoc(text=text, parts=split_text(text))

    pieces: List[str] = []
    parts: List[Dict] = []
    pos = 0
    for i, (title, body) in enumerate(nonempty):
        heading = title or f"第 {i + 1} 部分"
        if title and body.startswith(title):
            chunk = body
        else:
            chunk = f"{heading}\n\n{body}"
        if pieces:
            pos += 2  # 与 join 的 "\n\n" 对齐
        start = pos
        end = start + len(chunk)
        parts.append({"title": heading, "start": start, "end": end})
        pieces.append(chunk)
        pos = end
    text = "\n\n".join(pieces)
    return ExtractedDoc(text=text, parts=parts)


# ---------- EPUB ----------

def _extract_epub(data: bytes) -> ExtractedDoc:
    from ebooklib import ITEM_DOCUMENT, epub

    try:
        book = epub.read_epub(io.BytesIO(data), options={"ignore_ncx": True})
    except Exception:
        book = epub.read_epub(io.BytesIO(data))

    toc_titles: Dict[str, str] = {}

    def walk(nodes) -> None:
        for n in nodes or []:
            if isinstance(n, tuple) and n:
                walk([n[0]])
                if len(n) > 1:
                    walk(n[1])
            elif isinstance(n, epub.Link):
                href = (n.href or "").split("#")[0]
                if href and n.title:
                    toc_titles[href] = n.title
                    toc_titles[href.split("/")[-1]] = n.title
            elif isinstance(n, epub.Section):
                if getattr(n, "title", None) and getattr(n, "href", None):
                    href = (n.href or "").split("#")[0]
                    toc_titles[href] = n.title

    try:
        walk(book.toc)
    except Exception:
        pass

    chunks: List[tuple[Optional[str], str]] = []
    seen = set()
    spine_ids = [idref for idref, _ in (book.spine or [])]
    items = []
    for idref in spine_ids:
        item = book.get_item_with_id(idref)
        if item is not None:
            items.append(item)
    if not items:
        items = list(book.get_items_of_type(ITEM_DOCUMENT))

    for item in items:
        name = (item.get_name() or "").replace("\\", "/")
        base = name.split("/")[-1].lower()
        props = {str(p).lower() for p in (getattr(item, "properties", None) or [])}
        if "nav" in props or base in {"nav.xhtml", "nav.html", "toc.xhtml", "toc.html"}:
            continue
        if name in seen:
            continue
        seen.add(name)
        try:
            raw = item.get_content()
            html = raw.decode("utf-8") if isinstance(raw, (bytes, bytearray)) else str(raw)
        except Exception:
            continue
        body = _html_to_text(html)
        if not body:
            continue
        title = toc_titles.get(name) or toc_titles.get(name.split("/")[-1])
        if not title:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(html, "lxml")
            h = soup.find(["h1", "h2", "h3", "title"])
            if h:
                title = h.get_text(" ", strip=True) or None
        chunks.append((title, body))

    return _parts_from_chunks(chunks)


# ---------- DOCX / ODT / FB2 / RTF ----------

def _extract_docx(data: bytes) -> ExtractedDoc:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        raise ExtractError(f"无法解析 DOCX: {e}") from e
    lines: List[str] = []
    for p in doc.paragraphs:
        lines.append(p.text or "")
    for table in doc.tables:
        for row in table.rows:
            lines.append("\t".join((c.text or "").strip() for c in row.cells))
    from .chapters import split_text

    text = "\n".join(lines).strip()
    if not text:
        raise ExtractError("DOCX 无正文")
    return ExtractedDoc(text=text, parts=split_text(text))


def _extract_odt(data: bytes) -> ExtractedDoc:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            xml = zf.read("content.xml")
    except Exception as e:
        raise ExtractError(f"无法解析 ODT: {e}") from e
    try:
        root = ET.fromstring(xml)
    except ET.ParseError as e:
        raise ExtractError(f"ODT content.xml 损坏: {e}") from e
    paras = []
    for p in root.iter("{urn:oasis:names:tc:opendocument:xmlns:text:1.0}p"):
        paras.append("".join(p.itertext()))
    if not paras:
        # 无命名空间兜底
        paras = ["".join(p.itertext()) for p in root.iter() if p.tag.endswith("}p") or p.tag == "p"]
    from .chapters import split_text

    text = "\n".join(paras).strip()
    if not text:
        raise ExtractError("ODT 无正文")
    return ExtractedDoc(text=text, parts=split_text(text))


def _extract_fb2(data: bytes) -> ExtractedDoc:
    import re

    from .chapters import decode_bytes

    try:
        root = ET.fromstring(data)
    except ET.ParseError:
        # 非 UTF-8 的 FB2（常见 windows-1251）：先按中文字符集策略解码，
        # 再去掉 XML 声明里的 encoding（否则与重编码后的 UTF-8 冲突，仍会解析失败）
        text = re.sub(r"<\?xml[^?]*\?>", '<?xml version="1.0"?>', decode_bytes(data), count=1)
        root = ET.fromstring(text.encode("utf-8"))
    chunks: List[tuple[Optional[str], str]] = []
    for sec in root.iter():
        tag = sec.tag.rsplit("}", 1)[-1]
        if tag != "section":
            continue
        title_el = None
        for child in list(sec):
            if child.tag.rsplit("}", 1)[-1] == "title":
                title_el = child
                break
        title = "".join(title_el.itertext()).strip() if title_el is not None else None
        paras = []
        for child in sec:
            ct = child.tag.rsplit("}", 1)[-1]
            if ct in {"p", "poem", "cite", "subtitle"}:
                paras.append("".join(child.itertext()).strip())
        body = "\n".join(p for p in paras if p)
        if body:
            chunks.append((title, body))
    from .chapters import split_text

    if chunks:
        return _parts_from_chunks(chunks)
    text = _html_to_text(decode_bytes(data))
    if not text:
        raise ExtractError("FB2 无正文")
    return ExtractedDoc(text=text, parts=split_text(text))


def _extract_rtf(data: bytes) -> ExtractedDoc:
    from striprtf.striprtf import rtf_to_text

    raw = data.decode("latin-1", errors="ignore")
    try:
        text = rtf_to_text(raw)
    except Exception as e:
        raise ExtractError(f"无法解析 RTF: {e}") from e
    from .chapters import split_text

    text = (text or "").strip()
    if not text:
        raise ExtractError("RTF 无正文")
    return ExtractedDoc(text=text, parts=split_text(text))


# ---------- DOC（Word 97–2003 OLE，MS-DOC piece table） ----------

def _extract_doc(data: bytes) -> ExtractedDoc:
    import olefile

    try:
        ole = olefile.OleFileIO(io.BytesIO(data))
    except Exception as e:
        raise ExtractError(f"不是有效的 .doc: {e}") from e
    try:
        if not ole.exists("WordDocument"):
            raise ExtractError("缺少 WordDocument 流")
        wd = ole.openstream("WordDocument").read()
        text = _doc_piece_table_text(ole, wd) or _doc_scrape_text(wd)
    finally:
        ole.close()
    from .chapters import split_text

    text = (text or "").strip()
    if not text:
        raise ExtractError("DOC 无正文（可能是加密文档）")
    return ExtractedDoc(text=text, parts=split_text(text))


def _doc_piece_table_text(ole, wd: bytes) -> str:
    if len(wd) < 0x20:
        return ""
    magic = struct.unpack_from("<H", wd, 0)[0]
    if magic != 0xA5EC:
        return ""
    try:
        csw = struct.unpack_from("<H", wd, 32)[0]
        p = 34 + 2 * csw
        cslw = struct.unpack_from("<H", wd, p)[0]
        p += 2 + 4 * cslw
        n_pairs = struct.unpack_from("<H", wd, p)[0]
        p += 2
        fc_clx_index = 33
        if fc_clx_index >= n_pairs or p + 8 * (fc_clx_index + 1) > len(wd):
            return ""
        fc_clx, lcb_clx = struct.unpack_from("<II", wd, p + 8 * fc_clx_index)
    except struct.error:
        return ""
    if not lcb_clx:
        return ""
    flags = struct.unpack_from("<H", wd, 0x0A)[0]
    table_name = "1Table" if flags & 0x0200 else "0Table"
    if not ole.exists(table_name):
        return ""
    table = ole.openstream(table_name).read()
    if fc_clx + lcb_clx > len(table):
        return ""
    clx = table[fc_clx : fc_clx + lcb_clx]
    return _parse_clx(clx, wd)


def _parse_clx(clx: bytes, word_doc: bytes) -> str:
    i = 0
    n = len(clx)
    while i < n:
        tag = clx[i]
        if tag == 0x01:
            i += 1
            if i + 2 > n:
                break
            cb = struct.unpack_from("<h", clx, i)[0]
            i += 2 + max(cb, 0)
        elif tag == 0x02:
            i += 1
            if i + 4 > n:
                break
            lcb = struct.unpack_from("<I", clx, i)[0]
            i += 4
            return _parse_pcdt(clx[i : i + lcb], word_doc)
        else:
            break
    return ""


def _parse_pcdt(pcdt: bytes, word_doc: bytes) -> str:
    if len(pcdt) < 16:
        return ""
    # 4*(n+1) + 8*n = lcb  →  n = (lcb-4)/12
    if (len(pcdt) - 4) % 12 != 0:
        return ""
    n = (len(pcdt) - 4) // 12
    if n <= 0:
        return ""
    cps = [struct.unpack_from("<I", pcdt, 4 * i)[0] for i in range(n + 1)]
    pcd_off = 4 * (n + 1)
    out: List[str] = []
    for i in range(n):
        char_count = cps[i + 1] - cps[i]
        if char_count <= 0:
            continue
        fc_raw = struct.unpack_from("<I", pcdt, pcd_off + 8 * i + 2)[0]
        compressed = bool(fc_raw & 0x40000000)
        fc = fc_raw & 0x3FFFFFFF
        if compressed:
            fc = fc // 2
            end = fc + char_count
            if end > len(word_doc) or fc < 0:
                continue
            out.append(_decode_ansi(word_doc[fc:end]))
        else:
            end = fc + char_count * 2
            if end > len(word_doc) or fc < 0:
                continue
            out.append(word_doc[fc:end].decode("utf-16le", errors="replace"))
    text = "".join(out)
    # Word 用 \r 分段、\x07 单元格、\x0c 分页
    text = text.replace("\r", "\n").replace("\x07", "\t").replace("\x0c", "\n\n")
    text = text.replace("\x00", "")
    return text.strip()


def _decode_ansi(blob: bytes) -> str:
    for enc in ("gb18030", "cp1252"):
        try:
            return blob.decode(enc)
        except UnicodeDecodeError:
            continue
    return blob.decode("gb18030", errors="replace")


def _doc_scrape_text(wd: bytes) -> str:
    """piece table 失败时：从 UTF-16LE 里捞可读片段（中英文小说够用）。"""
    import re

    text = wd.decode("utf-16le", errors="ignore")
    runs = re.findall(
        r"[\u4e00-\u9fffA-Za-z0-9，。！？、；：“”‘’…—《》\s]{12,}",
        text,
    )
    return "\n".join(r.strip() for r in runs if r.strip())
