"""站内阅读：从 epub / docx / html / rtf / odt 抽出正文。"""
from __future__ import annotations

import io
import unittest
import zipfile

from app.extract import extract_document


def _make_epub() -> bytes:
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("gh-test")
    book.set_title("测试书")
    book.set_language("zh")
    c1 = epub.EpubHtml(title="第一章 相遇", file_name="c1.xhtml", lang="zh")
    c1.content = "<html><body><h1>第一章 相遇</h1><p>正文甲甲甲</p></body></html>"
    c2 = epub.EpubHtml(title="第二章 分别", file_name="c2.xhtml", lang="zh")
    c2.content = "<html><body><h1>第二章 分别</h1><p>正文乙乙乙</p></body></html>"
    book.add_item(c1)
    book.add_item(c2)
    book.toc = (c1, c2)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", c1, c2]
    buf = io.BytesIO()
    epub.write_epub(buf, book)
    return buf.getvalue()


def _make_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("第一章 开端")
    doc.add_paragraph("这是第一章的正文。")
    doc.add_paragraph("第二章 高潮")
    doc.add_paragraph("这是第二章的正文。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_odt() -> bytes:
    xml = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">
  <office:body><office:text>
    <text:p>第一章 测试</text:p>
    <text:p>ODT 正文内容</text:p>
  </office:text></office:body>
</office:document-content>
"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        zf.writestr("content.xml", xml)
    return buf.getvalue()


class ExtractTests(unittest.TestCase):
    def test_txt_split(self):
        raw = "序\n\n第一章 相遇\n甲\n\n第二章 分别\n乙".encode("utf-8")
        doc = extract_document(raw, "txt")
        self.assertGreaterEqual(len(doc.parts), 2)
        self.assertIn("甲", doc.text)

    def test_epub_spine_chapters(self):
        doc = extract_document(_make_epub(), "epub")
        self.assertIn("正文甲甲甲", doc.text)
        self.assertIn("正文乙乙乙", doc.text)
        self.assertGreaterEqual(len(doc.parts), 2)

    def test_docx_and_heading_split(self):
        doc = extract_document(_make_docx(), "docx")
        self.assertIn("第一章的正文", doc.text)
        self.assertGreaterEqual(len(doc.parts), 2)

    def test_html(self):
        html = "<html><body><h1>第一章</h1><p>你好</p><script>alert(1)</script></body></html>"
        doc = extract_document(html.encode("utf-8"), "html")
        self.assertIn("你好", doc.text)
        self.assertNotIn("alert", doc.text)

    def test_rtf(self):
        rtf = r"{\rtf1\ansi\deff0{\fonttbl{\f0 Times;}}\f0\fs24 hello world\par}"
        doc = extract_document(rtf.encode("ascii"), "rtf")
        self.assertIn("hello", doc.text.lower())

    def test_odt(self):
        doc = extract_document(_make_odt(), "odt")
        self.assertIn("ODT 正文内容", doc.text)

    def test_fb2_sections(self):
        xml = (
            '<?xml version="1.0" encoding="utf-8"?>'
            '<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0">'
            "<body>"
            '<section><title>第一章 相遇</title><p>正文内容甲</p></section>'
            '<section><title>第二章 分别</title><p>正文内容乙</p></section>'
            "</body></FictionBook>"
        )
        doc = extract_document(xml.encode("utf-8"), "fb2")
        self.assertIn("正文内容甲", doc.text)
        self.assertIn("正文内容乙", doc.text)
        self.assertGreaterEqual(len(doc.parts), 2)

    def test_fb2_misdeclared_encoding_fallback(self):
        # 声明 UTF-8 但实际是 GBK 字节（中文老书常见）：走编码兜底路径，不得抛 NameError
        xml = (
            '<?xml version="1.0" encoding="utf-8"?>'
            '<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0">'
            '<body><section><title>第一章</title><p>你好，世界</p></section></body>'
            "</FictionBook>"
        )
        doc = extract_document(xml.encode("gbk"), "fb2")
        self.assertIn("你好，世界", doc.text)

    def test_unknown_ext(self):
        with self.assertRaises(Exception):
            extract_document(b"xx", "mobi")


if __name__ == "__main__":
    unittest.main()
